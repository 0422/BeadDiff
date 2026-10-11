# 2026-10-10：本次修改原因：按妙手模板的 22 个字段导出已审核商品，记录导出使用的商品版本。
from __future__ import annotations

from pathlib import Path
from uuid import uuid4

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from .models import Product
from .storage import Repository

# 2026-10-11：本次修改原因：本地生成素材可编辑审核，但不能把本机文件地址交给远程 ERP。
from .content import description_images


HEADERS = [
    "*产品主编号", "*产品名称", "货币类型", "产品主图", "货源链接", "货源平台", "货源ID",
    "详情描述", "详情图", "货源类目", "自定义属性", "产品视频", "产品证书", "尺寸图表",
    "SKU规格1", "SKU规格2", "平台SKU", "*SKU售价", "SKU图片", "SKU库存", "SKU重量(KG)", "SKU尺寸(CM)",
]


def product_rows(product: Product) -> list[list[str]]:
    rows = []
    for index, sku in enumerate(product.skus):
        rows.append([
            product.main_code, product.title, product.currency.upper(),
            "，".join(product.images) if index == 0 else "", product.source_url, product.source, product.source_id,
            # product.description if index == 0 else "", "，".join(product.detail_images) if index == 0 else "",
            # 2026-10-11：本次修改原因：详情描述优先导出图文 HTML，旧商品人工清空后不恢复摘要。
            (product.description_html if product.description_html or "description_html" in product.manual_fields
             else product.description) if index == 0 else "", "，".join(product.detail_images) if index == 0 else "",
            product.category, "；".join(f"{key}:{value}" for key, value in product.attributes.items()),
            product.video, product.certificate, product.size_chart,
            sku.spec1, sku.spec2, sku.platform_sku, sku.price, sku.image, sku.stock, sku.weight, sku.dimensions,
        ])
    return rows


def export_products(repository: Repository, product_ids: list[str], path: Path) -> tuple[int, int]:
    products = [repository.product(product_id) for product_id in dict.fromkeys(product_ids)]
    if not products or any(product is None for product in products):
        raise ValueError("请选择仍在商品库中的商品。")
    codes = set()
    for product in products:
        if not product.approved or product.issues():
            raise ValueError(f"“{product.title or product.source_id}”尚未完善并审核，不能导出。")
        if product.main_code in codes:
            raise ValueError(f"产品主编号重复：{product.main_code}，请先修改。")
        codes.add(product.main_code)
        # 2026-10-11：本次修改原因：导出前准确提示需要转为公开链接的本地素材，不静默输出无效地址。
        media = product.images + product.detail_images + [sku.image for sku in product.skus]
        media += description_images(product.description_html)
        if any(address.startswith("file:") for address in media):
            raise ValueError(f"“{product.title or product.main_code}”含本地图片。请先上传至可公开访问的图片存储并替换地址，再导出妙手表格。")
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "模版"
    sheet.append(HEADERS)
    for product in products:
        for values in product_rows(product):
            # 2026-10-10：本次修改原因：阻止超长文本或非法控制字符被 Excel 静默截断。
            for column, value in zip(HEADERS, values):
                if len(value) > 32767:
                    raise ValueError(f"{product.main_code} 的“{column}”超过 Excel 单元格的 32767 字符限制。")
                if any(ord(character) < 32 and character not in "\t\n\r" for character in value):
                    raise ValueError(f"{product.main_code} 的“{column}”含 Excel 不支持的控制字符。")
            # 设为文本单元格，保留长 SKU/货源 ID，并避免商品文案被 Excel 当作公式。
            sheet.append(values)
            for cell in sheet[sheet.max_row]:
                cell.data_type = "s"
                cell.number_format = "@"
                cell.alignment = Alignment(vertical="top", wrap_text=False)
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="123D3B")
        cell.font = Font(name="Microsoft YaHei", color="FFFFFF", bold=True)
        cell.alignment = Alignment(vertical="center")
        sheet.column_dimensions[cell.column_letter].width = 24
    sheet.column_dimensions["B"].width = 52
    sheet.row_dimensions[1].height = 30
    sheet.freeze_panes = "C2"
    sheet.auto_filter.ref = sheet.dimensions
    path = path.with_suffix(".xlsx")
    # 2026-10-10：本次修改原因：避免采集材料或手工编辑出的编号被用作数据目录里的覆盖目标。
    if path.resolve().is_relative_to(repository.data_dir.resolve()):
        raise ValueError("请将导出文件保存到 data 目录之外，保持数据目录只保存数据库和采集材料。")
    temporary = path.with_name(f".{path.stem}-{uuid4().hex}.tmp.xlsx")
    try:
        workbook.save(temporary)
        temporary.replace(path)
    finally:
        workbook.close()
        temporary.unlink(missing_ok=True)
    repository.record_export(path, products)
    return len(products), sum(len(product.skus) for product in products)
