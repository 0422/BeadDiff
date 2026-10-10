# 2026-10-10：本次修改原因：为解析保守性、重复采集保护、审核失效和 22 列导出提供回归覆盖。
import json
from pathlib import Path
from urllib.parse import urlencode

import pytest
from openpyxl import load_workbook

from beaddiff.exporter import HEADERS, export_products
from beaddiff.parsing import parse_capture
from beaddiff.service import CaptureService, envelopes_from_file
from beaddiff.storage import Repository


def payload(**overrides):
    goods = {
        "goodsId": "TEST-1", "goodsName": "结构测试商品", "currency": "USD",
        "gallery": [{"url": "https://example.com/main.jpg"}],
        "skuList": [
            {"skuId": "BLACK", "price": 499, "specs": [{"parentSpecName": "颜色", "specName": "黑"}]},
            {"skuId": "RED", "priceStr": "$5.99", "specs": [{"parentSpecName": "颜色", "specName": "红"}]},
        ],
    }
    result = {"source": "temu", "itemId": "TEST-1", "site": "US", "afterUrl": "https://example.com/product",
              "pageContent": "<script>window.rawData = " + json.dumps({"store": {"goods": goods}}, ensure_ascii=False) + ";</script>"}
    result.update(overrides)
    return result


def test_prices_and_stock_are_not_invented():
    product = parse_capture(payload())
    assert product.title == "结构测试商品"
    assert product.spec1_name == "颜色"
    assert [sku.spec1 for sku in product.skus] == ["黑", "红"]
    assert product.skus[0].source_price == "499"
    assert all(sku.price == "" and sku.stock == "" and sku.weight == "" for sku in product.skus)


def test_unrelated_recommended_goods_are_not_adopted():
    product = parse_capture(payload(itemId="DIFFERENT"))
    assert not product.title and not product.skus
    assert any("不匹配" in warning for warning in product.warnings)


def test_missing_data_does_not_create_sku_or_currency():
    product = parse_capture(payload(pageContent='<meta property="og:title" content="摘要商品">'))
    assert product.title == "摘要商品"
    assert product.skus == []
    assert product.currency == ""
    with pytest.raises(ValueError, match="itemId"):
        parse_capture({"pageContent": "", "source": "temu"})


# 2026-10-10：本次修改原因：回归覆盖超出妙手两维规格限制及压缩材料保守处理。
def test_extra_spec_dimension_is_preserved_and_blocks_approval():
    material = payload()
    material["rawData"] = {"store": {"goods": {
        "goodsId": "TEST-1", "goodsName": "多维规格", "currency": "USD",
        "skuList": [{"skuId": "THREE", "specs": {"颜色": "黑", "尺寸": "8mm", "包装": "一袋"}}],
    }}}
    material["pageContent"] = ""
    product = parse_capture(material)
    assert len(product.skus[0].specs) == 3
    assert any("超过两个规格" in issue for issue in product.issues())
    compressed = parse_capture(payload(isZip="true"))
    assert compressed.skus == []
    assert any("压缩" in warning for warning in compressed.warnings)


def test_retry_and_recapture_preserve_manual_edits(tmp_path: Path):
    repository = Repository(tmp_path)
    service = CaptureService(repository)
    first = service.ingest({"eventId": "one", "payload": payload()})
    retry = service.ingest({"eventId": "one", "payload": payload()})
    assert retry["duplicate"]
    assert len(repository.captures()) == 1
    product = repository.product(first["productId"])
    product.title = "人工修改标题"
    repository.save_product(product)
    recapture = service.ingest({"eventId": "two", "payload": payload()})
    assert recapture["state"] == "重复商品"
    assert len(repository.captures()) == 2
    assert repository.product(product.id).title == "人工修改标题"


def test_approval_invalidated_and_stale_edit_rejected(tmp_path: Path):
    repository = Repository(tmp_path)
    service = CaptureService(repository)
    result = service.ingest(payload())
    product = repository.product(result["productId"])
    with pytest.raises(ValueError, match="完善"):
        repository.save_product(product, approve=True)
    for sku in product.skus:
        sku.price = "9.90"
    approved = repository.save_product(product, approve=True)
    assert approved.approved
    stale = repository.product(product.id)
    approved.title = "修改后待重新审核"
    saved = repository.save_product(approved)
    assert not saved.approved
    with pytest.raises(ValueError, match="其他操作"):
        repository.save_product(stale)


def test_har_and_urlencoded_import_extract_body_only(tmp_path: Path):
    form = payload()
    txt = tmp_path / "capture.txt"
    txt.write_text(urlencode(form), encoding="utf-8")
    assert envelopes_from_file(txt)[0]["payload"] == form
    har = tmp_path / "capture.har"
    har.write_text(json.dumps({"log": {"entries": [{"request": {
        "url": "https://kuajing-plw.happyporsche.com/open/fetch/pfti",
        "headers": [{"name": "Cookie", "value": "do-not-copy"}],
        "postData": {"mimeType": "application/x-www-form-urlencoded", "text": urlencode(form)},
    }}]}}), encoding="utf-8")
    envelope = envelopes_from_file(har)[0]
    assert envelope["payload"] == form
    assert "do-not-copy" not in json.dumps(envelope)


def test_export_requires_approval_and_keeps_ids_as_text(tmp_path: Path):
    repository = Repository(tmp_path / "data")
    result = CaptureService(repository).ingest(payload())
    product = repository.product(result["productId"])
    path = tmp_path / "export.xlsx"
    with pytest.raises(ValueError, match="审核"):
        export_products(repository, [product.id], path)
    product.title = "=1+1"
    product.skus[0].platform_sku = "00000000000000000001"
    for sku in product.skus:
        sku.price = "12.50"
    repository.save_product(product, approve=True)
    assert export_products(repository, [product.id], path) == (1, 2)
    book = load_workbook(path)
    sheet = book.active
    assert [cell.value for cell in sheet[1]] == HEADERS
    assert sheet.max_column == 22 and sheet.max_row == 3
    assert sheet["Q2"].value == "00000000000000000001"
    assert sheet["B2"].value == "=1+1" and sheet["B2"].data_type == "s"
    assert sheet["D2"].value and sheet["D3"].value is None
    assert len(repository.exports()) == 1
    book.close()
